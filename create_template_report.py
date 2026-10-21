from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

TEMPLATE = Path('/home/ubuntu/upload/Internship_Report.docx')
OUT = Path('/home/ubuntu/HARVIX---AIBOT/Internship_Report_HARVIX_Template.docx')
SCREENSHOT = Path('/home/ubuntu/HARVIX---AIBOT/docs/dashboard-screenshot.png')
NAVY='17365D'; BLUE='2F75B5'; GRAY='666666'

doc = Document(TEMPLATE)

# Remove everything from the existing Chapter 1 onward, preserving all template content before it.
chapter_start = None
for p in doc.paragraphs:
    if p.text.strip() == 'CHAPTER 1':
        chapter_start = p._p
        break
if chapter_start is None:
    raise RuntimeError('CHAPTER 1 was not found in the supplied template')
body = doc._element.body
children = list(body)
start_index = children.index(chapter_start)
for child in children[start_index:]:
    if child.tag == qn('w:sectPr'):
        break
    body.remove(child)

# Page-number footer is retained from the template; ensure it remains visible.
for section in doc.sections:
    for p in section.footer.paragraphs:
        for r in p.runs:
            r.font.name='Calibri'; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(GRAY)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell(cell, text, bold=False, color=None, size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(str(text)); r.font.name='Calibri'; r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def heading(text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.space_before=Pt(14 if level==1 else 9); p.paragraph_format.space_after=Pt(5); p.add_run(text); return p

def body_text(text):
    p=doc.add_paragraph(style='Body Text'); p.paragraph_format.space_after=Pt(7); p.paragraph_format.line_spacing=1.15; p.add_run(text); return p

def bullet(text):
    p=doc.add_paragraph(style='List Paragraph'); p.paragraph_format.left_indent=Inches(0.25); p.paragraph_format.first_line_indent=Inches(-0.18); p.paragraph_format.space_after=Pt(4); p.add_run('• ' + text); return p

def numbered(text):
    p=doc.add_paragraph(style='List Paragraph'); p.paragraph_format.left_indent=Inches(0.25); p.paragraph_format.first_line_indent=Inches(-0.18); p.paragraph_format.space_after=Pt(4); p.add_run(text); return p

def caption(text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8); r=p.add_run(text); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(GRAY); return p

def table(headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    if 'Table Grid' in [s.name for s in doc.styles]:
        t.style='Table Grid'
    repeat_header(t.rows[0])
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i],h,True,'FFFFFF'); shade(t.rows[0].cells[i],NAVY)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row): set_cell(cells[i],v)
        if ri%2==0:
            for c in cells: shade(c,'F4F8FB')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def code_block(text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25); p.paragraph_format.space_after=Pt(5); r=p.add_run(text); r.font.name='Consolas'; r.font.size=Pt(9); return p

# Start Chapter 1 on a new page after the preserved template pages.
doc.add_page_break()

# CHAPTER 1
heading('CHAPTER 1',1); heading('Introduction',2)
heading('Background',5)
body_text('HARVIX is an AI-powered web assistant developed to support users with programming, mathematics, and general queries through a single conversational platform. The system provides programming assistance such as code generation, debugging, code explanation, refactoring, and code conversion. It also supports mathematical problems including equations, derivatives, integrals, and symbolic expressions.')
body_text('The application uses React and Vite for the frontend and FastAPI for the backend. SymPy is integrated to perform exact mathematical calculations instead of depending only on language-model predictions. An optional OpenAI-compatible API can be configured to provide richer conversational and coding responses, while a local fallback keeps the project usable during development and demonstration.')
heading('Problem Statement',5)
body_text('Students, developers, and technical learners often use separate tools for coding assistance, mathematical calculations, and general explanations. This process can be time-consuming and inconvenient. Search engines and documentation require users to find, compare, and understand information on their own, while general AI tools may produce incorrect code or inaccurate mathematical answers. HARVIX addresses this need by identifying the request type and routing it to a suitable processing method.')
heading('Motivation',5)
body_text('The motivation behind HARVIX is to provide quick and convenient support for students, programmers, and learners. The system reduces the time required to search for programming solutions and mathematical explanations by bringing multiple capabilities into one workspace. SymPy improves the reliability of mathematical results, while the optional AI integration supports natural-language explanations and coding assistance. The project also demonstrates practical use of frontend development, backend API design, artificial intelligence integration, symbolic computation, testing, Docker, and continuous integration.')
heading('Objectives of the Project',5)
for x in ['To develop a web-based AI assistant named HARVIX.', 'To create a responsive and user-friendly dashboard using React and Vite.', 'To develop a REST API using Python FastAPI.', 'To automatically classify requests into general, coding, and mathematics categories.', 'To provide code generation, debugging, explanation, refactoring, and code conversion features.', 'To support equations, derivatives, integrals, and symbolic simplification.', 'To integrate SymPy for accurate and exact mathematical calculations.', 'To provide optional OpenAI-compatible language-model integration.', 'To implement input validation, error handling, health monitoring, automated tests, Docker, and GitHub Actions.']: bullet(x)
heading('Scope of the Project',5)
body_text('The project covers the design and development of a responsive conversational dashboard, FastAPI backend, automatic intent-routing system, SymPy-based mathematics engine, optional AI provider integration, testing framework, Docker configuration, GitHub Actions workflow, and project documentation. The current version supports general conversation, programming assistance, mathematical problem solving, and responsive web access. Persistent authentication, server-side chat history, real-time streaming, database integration, and secure execution of untrusted code are outside the current implementation and are reserved for future development.')
heading('Organization of the Report',5)
body_text('Chapter 2 reviews existing approaches and their limitations. Chapter 3 presents the proposed HARVIX architecture and functional modules. Chapter 4 explains the development methodology, technology stack, request routing, mathematics processing, and testing approach. Chapter 5 describes the implementation, project structure, API summary, challenges, and validation. Chapter 6 concludes the report with achievements, results, limitations, future scope, and learning outcomes.')

# CHAPTER 2
heading('CHAPTER 2',1); heading('Existing Method',2)
heading('Overview',5); body_text('Before developing HARVIX, common approaches for obtaining programming and mathematics assistance were considered. Each method provides value but also has limitations related to availability, flexibility, accuracy, or verification.')
heading('Human-Staffed Support',5); body_text('Teachers, mentors, and experienced developers can provide context-aware guidance, but their availability is limited. Human support does not scale easily when many users ask repetitive questions, and responses may be delayed during busy periods.')
heading('Static FAQ Pages and Keyword Search',5); body_text('Search engines and documentation provide broad technical information. However, users must select appropriate keywords, compare multiple sources, and adapt the information to their own problem. Beginners may also find it difficult to identify which result is correct or current.')
heading('Rule-Based Chatbots',5); body_text('Rule-based chatbots respond according to manually defined patterns and decision trees. They are predictable but have narrow coverage. A request with unfamiliar wording or multiple requirements may not match an existing rule and may receive an unhelpful fallback response.')
heading('Plain LLM Chatbots Without Retrieval',5); body_text('General language-model chatbots can generate natural explanations and code, but they may produce incorrect code or unsupported calculations. Mathematical responses are especially important to verify because a fluent answer is not proof of a correct result.')
heading('Summary of Limitations',5)
for x in ['Human assistance is limited by time, availability, and scalability.', 'Search and documentation require users to perform the interpretation and verification themselves.', 'Rule-based chatbots have limited flexibility and require continuous manual rule creation.', 'General AI chatbots may hallucinate and should not be the only source for exact mathematical calculations.']: bullet(x)

# CHAPTER 3
heading('CHAPTER 3',1); heading('Proposed Method with Architecture',2)
heading('Overview of the Proposed Solution',5)
body_text('HARVIX combines a responsive conversation dashboard with a FastAPI service and a tool-aware request router. When a user submits a message, the system determines whether the message is a general question, coding request, or mathematics request. Mathematics requests are sent to the SymPy engine, while general and coding requests use the local response layer or an optional OpenAI-compatible provider.')
heading('Why RAG over the Alternatives',5)
body_text('The template heading is retained for structural consistency. HARVIX does not currently implement a document-retrieval pipeline; instead, it focuses on direct intent routing and deterministic mathematics. A future retrieval layer could add project documentation or organizational knowledge as context before generation, while preserving the same principle of reducing unsupported answers.')
heading('System Architecture',5)
code_block('React/Vite Dashboard  →  FastAPI REST API  →  Intent Router\n                                                     ├── SymPy Math Engine\n                                                     ├── Optional OpenAI Provider\n                                                     └── Local Deterministic Fallback')
caption('Figure 3.1: HARVIX system architecture')
heading('Component Responsibilities',5)
body_text('The React frontend is responsible for the user interface, message composition, mode controls, suggestions, conversation display, and responsive layout. The FastAPI backend validates requests, exposes REST endpoints, routes messages, calls the mathematics engine, and optionally communicates with an external language-model provider. SymPy performs exact symbolic operations. Docker and GitHub Actions support repeatable execution and automated validation.')
heading('React Frontend',7); body_text('The React/Vite dashboard provides the HARVIX workspace, sidebar navigation, mode controls, suggested prompts, message bubbles, typing feedback, copy functionality, and responsive mobile behavior.')
heading('FastAPI Backend',7); body_text('The FastAPI service provides health, chat, and mathematics endpoints. It validates request fields, detects intent, returns structured responses, and keeps optional AI-provider credentials on the server side.')
heading('MongoDB Atlas',7); body_text('MongoDB Atlas is not used in the current HARVIX implementation. Conversation history is maintained in the frontend session, while persistent database storage is identified as a future extension.')
heading('ChromaDB',7); body_text('ChromaDB is not used in the current implementation because HARVIX does not yet include document ingestion or vector retrieval. A vector store could be added in a future RAG extension.')
heading('Google Gemini API',7); body_text('The current project uses an optional OpenAI-compatible provider rather than Google Gemini. The heading is retained to match the supplied template, while the implemented provider configuration is documented accurately in the technology stack.')
heading('Role-Based Access Design',5)
body_text('The current HARVIX version is a single-workspace demonstration and does not yet implement user roles or authentication. Role-based access can be added when persistent users, saved conversations, and administrative controls are introduced.')
heading('Advantages of the Proposed System',5)
for x in ['Unified access to coding, mathematics, and general assistance.', 'More reliable mathematical results through symbolic computation.', 'Responsive dashboard suitable for desktop and mobile screens.', 'Local fallback for development without a mandatory external service.', 'Clear project structure that supports testing, Docker, and future cloud deployment.']: bullet(x)
heading('Functional Modules of the System',5)
heading('Functional Modules',7); body_text('The system is organized into conversational assistance, coding assistance, mathematics solving, API configuration, and validation modules. Each module has a focused responsibility so that later features can be added without rewriting the whole application.')
heading('User Authentication Module',7); body_text('Authentication is not included in the current release. Future versions can add email/password and Google login with JWT-based session management and user-specific conversation history.')
heading('Customer Module',7); body_text('The current dashboard represents the end-user workspace. Users can start a conversation, select Auto, Coding, or Math mode, submit prompts, view responses, and use suggested actions.')
heading('Administrator Module',7); body_text('An administrator dashboard is outside the current scope. Future administration features may include usage monitoring, configuration management, user management, and service-health reporting.')
heading('Knowledge Base Management Module',7); body_text('Document upload and knowledge-base indexing are not implemented in the current HARVIX release. They are suitable future additions if a retrieval-augmented generation workflow is required.')
heading('Analytics Module',7); body_text('The current API returns response timing for each chat request, but it does not persist analytics records. A future analytics module could store conversations, response times, usage counts, and feedback.')

# CHAPTER 4
heading('CHAPTER 4',1); heading('Methodology',2)
heading('Technology Stack',5)
table(['Layer','Technology','Purpose'], [('Frontend','React, Vite, CSS, Lucide icons','Responsive dashboard and interaction controls'),('Backend','Python, FastAPI, Uvicorn','REST API, routing, validation, and service layer'),('Math engine','SymPy','Exact symbolic simplification, equations, derivatives, and integrals'),('AI provider','OpenAI-compatible API (optional)','Natural-language coding and general responses'),('Testing','Pytest and FastAPI TestClient','Endpoint and routing verification'),('DevOps','Docker, Docker Compose, GitHub Actions','Repeatable execution and continuous integration')])
heading('Document Ingestion Methodology',5)
body_text('HARVIX does not currently ingest documents or build a vector index. The current methodology starts from a user message and routes it directly to the appropriate assistant or mathematics tool. Document ingestion can be added later if the platform is extended with a retrieval-augmented knowledge base.')
heading('Query-Time Retrieval and Generation Methodology',5)
body_text('The implemented query-time flow receives the user message at POST /api/chat, determines the request mode, and creates a local or optional provider response. There is no retrieval step in the current version. A future RAG flow could retrieve relevant documentation, add it to the prompt, and return citations with the generated response.')
heading('Authentication and Authorization Methodology',5)
body_text('Authentication and authorization are not implemented in the current single-workspace version. API keys are kept in backend environment variables, and the .env file is excluded from Git. Future JWT or OAuth authentication can protect user history and administrative routes.')
heading('Admin Knowledge Base Management Workflow',5)
body_text('The current project does not include an administrator knowledge-base workflow. Future versions may provide document upload, validation, indexing, rebuild, and status controls for administrators.')
heading('Support Ticket and Feedback Methodology',5)
body_text('Support tickets and response feedback are outside the current implementation. The present application focuses on direct conversational assistance; ticket creation, ratings, and feedback analytics can be added around the chat experience later.')
heading('Analytics Methodology',5)
body_text('Each chat response includes elapsed_ms, allowing immediate response-time observation. Persistent analytics are not currently stored. A future analytics service can record request volume, selected modes, response times, errors, and user feedback.')
heading('Grounding and Hallucination-Prevention Strategy',5)
body_text('HARVIX reduces calculation risk by routing mathematical requests to SymPy rather than asking a language model to guess. The system prompt instructs the optional provider to be concise, ask for clarification when needed, use calculation tools for mathematics, and never fabricate results. Full document grounding is future scope because the current release has no retrieval database.')
heading('Working of Retrieval-Augmented Generation (RAG)',5)
body_text('RAG is documented here as a future extension rather than a claim about the current implementation. A future HARVIX RAG workflow would load approved documents, split them into chunks, create embeddings, retrieve relevant chunks for a question, and pass the retrieved context to the language model before generating an answer.')
heading('Advantages of Retrieval-Augmented Generation',5)
body_text('If added in the future, RAG could provide organization-specific answers, reduce unsupported responses, support citations, and allow knowledge updates without retraining the language model. The present HARVIX release instead provides a smaller, deterministic architecture centered on intent routing and SymPy.')

# CHAPTER 5
heading('CHAPTER 5',1); heading('Implementation',2)
heading('Backend Project Structure',5)
body_text('The backend is intentionally compact and easy to understand. The main module contains the FastAPI application, request models, intent router, SymPy math engine, local responses, and optional language-model adapter. Tests are kept in a separate backend/tests directory.')
code_block('backend/\n├── __init__.py\n├── app/\n│   ├── __init__.py\n│   └── main.py\n└── tests/\n    └── test_api.py')
heading('Authentication Implementation',5)
body_text('Authentication is not implemented in the current version. The application is delivered as a core AI workspace demonstration. Future implementation can add secure password hashing, JWT sessions, Google login, and user-specific access to saved conversations.')
heading('Multi-Format Document Ingestion',5)
body_text('Multi-format document ingestion is not part of the current HARVIX implementation. The application currently processes user messages directly. PDF, DOCX, TXT, and Markdown ingestion can be added if a future knowledge-base and RAG module is required.')
heading('Chat Endpoint with Background Analytics Logging',5)
body_text('The POST /api/chat endpoint accepts a message, optional mode, and recent history. It detects the intent, creates a local response, and attempts optional provider generation when configured. The response returns reply, mode, elapsed_ms, and sources. Persistent background analytics logging is future work; elapsed time is returned immediately for basic monitoring.')
heading('Knowledge Base Upload and Rebuild',5)
body_text('Knowledge-base upload and rebuild are not implemented in the current release. The current system does not claim to index organizational documents or retrieve document chunks.')
heading('Frontend Route Protection',5)
body_text('The current dashboard is a single workspace and does not require login-based route protection. Future authenticated versions can add guest and protected route guards so that user and administrator areas are separated.')
heading('API Summary',5)
table(['Endpoint','Method','Purpose'], [('/api/health','GET','Confirm that the API service is available'),('/api/chat','POST','Route a message and return a HARVIX response'),('/api/math','POST','Evaluate a symbolic expression or operation')])
heading('Development Challenges Encountered and Resolved',5)
for x in ['The GitHub repository was initially an empty scaffold, so the complete frontend, backend, tests, Docker configuration, and documentation were created from the specification.', 'The initial test run required explicit pytest path configuration; pytest.ini was added so package imports work consistently in local and CI environments.', 'The GitHub README architecture diagram produced a rich-rendering chunk error; it was replaced with a GitHub-safe text diagram.', 'The supplied report template described an unrelated RAG application; the Chapter 1–6 content was rewritten to match the implemented HARVIX project while preserving the template headings.']: bullet(x)
heading('Testing and Validation',5)
body_text('The completed application was validated through automated tests, direct API smoke tests, a production frontend build, Git integrity checks, and GitHub Actions. The live dashboard was also opened and captured for the report. The final template-based report was rendered to PDF to verify that the preserved first five pages and new chapters display correctly.')
p=doc.add_paragraph(); p.add_run('Functional Test Cases')
table(['Test Case','Input / Action','Expected Output','Result'], [('Health check','GET /api/health','status = ok','Pass'),('Equation solver','2*x^2 - 5*x + 3 = 0','Roots [1, 3/2]','Pass'),('Coding route','Python coding request','mode = coding','Pass'),('Frontend build','npm run build','Production assets generated','Pass'),('CI workflow','Push to main','Automated test job succeeds','Pass')])

# CHAPTER 6
heading('CHAPTER 6',1); heading('Conclusion',2)
heading('Summary',5)
body_text('The HARVIX internship project delivered a functional AI coding and mathematics assistant with a responsive dashboard, FastAPI backend, exact symbolic mathematics support, optional language-model integration, automated tests, Docker configuration, CI, and documentation. The application demonstrates how a natural-language interface can be combined with deterministic tools to provide more dependable assistance for specialized tasks.')
heading('Key Achievements',5)
for x in ['Developed a responsive React/Vite dashboard with a polished dark-theme interface.', 'Implemented FastAPI endpoints for health, chat, and mathematics processing.', 'Created automatic routing for general, coding, and mathematics requests.', 'Integrated SymPy for exact symbolic calculations.', 'Added optional backend-only OpenAI-compatible integration with a local fallback.', 'Added tests, Docker support, environment configuration, CI, and documentation.', 'Published and verified the completed project on GitHub.']: bullet(x)
heading('Results',5)
body_text('The project passed the backend test suite with three automated tests, successfully built the frontend production bundle, returned valid responses from the health, math, and chat endpoints, and completed the GitHub Actions CI workflow. The dashboard rendered successfully and was included as a project screenshot.')
heading('Limitations',5)
for x in ['Conversation history is currently maintained in the frontend session rather than a persistent user database.', 'Email/password and Google authentication are not implemented in the current version.', 'The typing indicator is not yet a true token-by-token streaming response.', 'User-submitted code is not executed, so the system does not currently provide runtime output.', 'Advanced mathematics such as matrices, statistics, unit conversion, and limits can be expanded in future releases.']: bullet(x)
heading('Future Scope',5)
for x in ['Add PostgreSQL-backed users and conversation history with JWT and OAuth authentication.', 'Implement Server-Sent Events for real-time response streaming.', 'Add a strictly isolated Python and JavaScript code sandbox with resource limits.', 'Expand the mathematics tool layer for matrices, statistics, units, limits, and verification by substitution.', 'Add rate limiting, caching, structured logging, and cloud monitoring.', 'Improve Markdown, code syntax highlighting, and LaTeX rendering in assistant responses.', 'Deploy the frontend and backend to cloud hosting with HTTPS and production secrets management.']: bullet(x)
heading('Learning Outcomes',5)
body_text('The project provided practical experience in full-stack development, React interface design, FastAPI REST services, symbolic mathematics, AI provider integration, environment-based configuration, automated testing, Docker, GitHub Actions, and technical reporting. It also demonstrated the importance of using deterministic tools for exact calculations and clearly separating implemented functionality from future enhancements.')
heading('Real-World Applications',5)
body_text('HARVIX can support programming education, coding bootcamps, developer onboarding, technical support teams, internal engineering assistance, and self-learning platforms. It can also be adapted for mathematics tutoring, software documentation assistants, API support tools, and educational laboratories. With authentication, persistence, and secure execution added, the platform can be extended into a multi-user cloud service.')

# Metadata and save
doc.core_properties.title='HARVIX Internship Report – Template Based'
doc.core_properties.subject='Chapters 1–6 report for HARVIX AI Coding and Mathematics Assistant'
doc.core_properties.author='Mohaneshwari P'
doc.core_properties.keywords='HARVIX, internship, React, FastAPI, SymPy, AI assistant'
doc.save(OUT)
print(OUT)
